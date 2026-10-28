import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    NAVY = RGBColor(0x2C, 0x3E, 0x50)       # #2C3E50
    DARK_BG = RGBColor(0x1A, 0x25, 0x2F)    # #1A252F
    TEAL = RGBColor(0x1B, 0xBC, 0x9B)       # #1BBC9B
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    LIGHT_GRAY = RGBColor(0xF8, 0xF9, 0xFA)
    BORDER_GRAY = RGBColor(0xDC, 0xE4, 0xEC)
    TEXT_DARK = RGBColor(0x2D, 0x37, 0x48)
    TEXT_MUTED = RGBColor(0x71, 0x80, 0x96)
    
    blank_layout = prs.slide_layouts[6]
    
    def add_header(slide, title_text, step_badge="TUTORIAL PENGGUNAAN"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.5), Inches(0.3))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = step_badge.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL
        
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.5), Inches(0.65))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY
        
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_GRAY
        line.line.color.rgb = BORDER_GRAY

    def add_card(slide, left, top, width, height, title, lines, title_color=NAVY, bg_color=LIGHT_GRAY, border_color=BORDER_GRAY):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)
        
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = title_color
        
        for line in lines:
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK
            p.space_before = Pt(4)
        return card

    def add_image_card(slide, left, top, width, height, image_path, caption=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_GRAY
        card.line.width = Pt(1)
        
        img_top = top + Inches(0.15)
        img_height = height - (Inches(0.6) if caption else Inches(0.3))
        img_left = left + Inches(0.15)
        img_width = width - Inches(0.3)
        
        if os.path.exists(image_path):
            slide.shapes.add_picture(image_path, img_left, img_top, img_width, img_height)
            
        if caption:
            cap_box = slide.shapes.add_textbox(left + Inches(0.2), top + height - Inches(0.45), width - Inches(0.4), Inches(0.35))
            tf_cap = cap_box.text_frame
            p_cap = tf_cap.paragraphs[0]
            p_cap.text = caption
            p_cap.font.size = Pt(10)
            p_cap.font.italic = True
            p_cap.font.color.rgb = TEXT_MUTED
            p_cap.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 1: Judul Presentasi (TANPA GAMBAR / NO IMAGE)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BG
    bg1.line.fill.background()
    
    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.9), Inches(3.6), Inches(0.36))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = TEAL
    badge1.line.fill.background()
    p_b = badge1.text_frame.paragraphs[0]
    p_b.text = "PANDUAN PRAKTIKUM MANDIRI"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = WHITE
    p_b.alignment = PP_ALIGN.CENTER
    
    t_box1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(1.6))
    tf1 = t_box1.text_frame
    tf1.word_wrap = True
    p1_1 = tf1.paragraphs[0]
    p1_1.text = "Tutorial Penggunaan Mobile VPLAS Verilog"
    p1_1.font.size = Pt(36)
    p1_1.font.bold = True
    p1_1.font.color.rgb = WHITE
    
    p1_2 = tf1.add_paragraph()
    p1_2.text = "Panduan Praktis Pembelajaran Mandiri Desain Sirkuit Digital & Sintaks Verilog HDL"
    p1_2.font.size = Pt(18)
    p1_2.font.color.rgb = TEAL
    p1_2.space_before = Pt(4)
    
    # 2 Cards for 5 Core Steps Overview (Slide 1 without image)
    card_l1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.2), Inches(5.6), Inches(3.7))
    card_l1.fill.solid()
    card_l1.fill.fore_color.rgb = NAVY
    card_l1.line.color.rgb = TEAL
    tf_cl1 = card_l1.text_frame
    tf_cl1.word_wrap = True
    p_cl1 = tf_cl1.paragraphs[0]
    p_cl1.text = "Tahapan Eksplorasi Soal & Kode:"
    p_cl1.font.size = Pt(13)
    p_cl1.font.bold = True
    p_cl1.font.color.rgb = TEAL
    steps_left = [
        ("Bagian 1", "Halaman Utama: Problems List & Navigasi Soal"),
        ("Bagian 2", "Anatomi Soal: Diagram Sirkuit & Knowledge Point"),
        ("Bagian 3", "Pengisian Kode Verilog: Dropdown & Fitur Auto-Sizing"),
    ]
    for stp, desc in steps_left:
        p = tf_cl1.add_paragraph()
        p.text = f"• [{stp}]  {desc}"
        p.font.size = Pt(12)
        p.font.color.rgb = WHITE
        p.space_before = Pt(12)

    card_r1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(3.2), Inches(5.6), Inches(3.7))
    card_r1.fill.solid()
    card_r1.fill.fore_color.rgb = NAVY
    card_r1.line.color.rgb = TEAL
    tf_cr1 = card_r1.text_frame
    tf_cr1.word_wrap = True
    p_cr1 = tf_cr1.paragraphs[0]
    p_cr1.text = "Tahapan Pertanyaan & Ketuntasan:"
    p_cr1.font.size = Pt(13)
    p_cr1.font.bold = True
    p_cr1.font.color.rgb = TEAL
    steps_right = [
        ("Bagian 4", "Menjawab Pertanyaan: Murni Pilihan Ganda (Tanpa Text Box)"),
        ("Bagian 5", "Statistical Data: Evaluasi & Umpan Balik Real-Time"),
        ("Keunggulan", "100% Offline, Bebas Typo, Teks Utuh & Ramah Sentuhan"),
    ]
    for stp, desc in steps_right:
        p = tf_cr1.add_paragraph()
        p.text = f"• [{stp}]  {desc}"
        p.font.size = Pt(12)
        p.font.color.rgb = WHITE
        p.space_before = Pt(12)

    # =========================================================================
    # SLIDE 2: Problems List (Halaman Utama Daftar Soal & Navigasi)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Halaman Utama: Problems List & Navigasi Soal", "BAGIAN 1 DARI 5")
    add_card(s2, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Daftar 22 Soal Bab 1 & Status Pengerjaan Real-Time",
        [
            "Daftar 22 Soal Terstruktur Bab 1:",
            "  * Soal 1-5  : Port wire, continuous assign, modul NOT gate.",
            "  * Soal 6-11 : Gerbang logika AND, OR, XOR, NAND, NOR.",
            "  * Soal 12-16: Multiplexer 2-to-1, Decoder, Adder, Subtractor.",
            "  * Soal 17-22: Rangkaian sekuensial Flip-Flop, Counter, & FSM.",
            "",
            "Badge Remark Status di Sisi Kanan Baris:",
            "  * Not Attempted (Polos): Belum pernah dikerjakan.",
            "  * Tried (Kuning): Pernah dicoba namun belum sempurna.",
            "  * Completed (Hijau): Sempurna (100% Benar / All Correct).",
            "",
            "Bilah Progres Sesi & Jalan Pintas TOP:",
            "  * Kartu Progres di bagian atas menghitung capaian (misal: 22/22) secara real-time.",
            "  * Tombol 'TOP' di AppBar untuk jalan pintas kembali ke puncak daftar.",
            "  * Cukup ketuk baris soal untuk membuka layar pengerjaan."
        ]
    )
    add_image_card(s2, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_problems_list.png", "Tampilan Halaman Utama: Daftar 22 Soal & Bilah Progres Pengerjaan")

    # =========================================================================
    # SLIDE 3: Diagram & Teori Soal (Diagram Sirkuit Teratas & Knowledge Point)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Anatomi Soal: Diagram Sirkuit & Knowledge Point", "BAGIAN 2 DARI 5")
    add_card(s3, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Pemahaman Visual Skema Rangkaian & Konsep Teori",
        [
            "1. Circuit Diagrams & Waveforms di Posisi Paling Atas Layar:",
            "   * Gambar skema sirkuit dan gelombang timing selalu muncul di bagian teratas sebelum kode Verilog.",
            "   * Membantu mahasiswa menganalisis keterhubungan pin dan sinyal logika terlebih dahulu.",
            "",
            "2. Fitur Zoom Interaktif:",
            "   * Dilengkapi gestur cubit (Pinch to Zoom) dan double-tap untuk memperbesar detail sinyal maupun teks skema.",
            "",
            "3. Kartu Teori 'Knowledge Point':",
            "   * Merangkum konsep kunci Verilog penting yang sedang diuji pada soal tersebut (misal: aturan continuous assign, port wire, dll.).",
            "",
            "4. Editor Kode Verilog HDL:",
            "   * Blok modul rapi bertema gelap dengan penomoran baris untuk mempermudah identifikasi bagian rumpang."
        ]
    )
    add_image_card(s3, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "assets/images/Verilog_CH1_question2_img1.png", "Diagram Skema Rangkaian di Posisi Teratas Layar Soal")

    # =========================================================================
    # SLIDE 4: Pengisian Kode Verilog (Code Blanks & Auto-Sizing)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Tutorial Pengisian Kode: Dropdown & Auto-Sizing", "BAGIAN 3 DARI 5")
    add_card(s4, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Pengisian Isian Kode Verilog Cerdas & Teks Utuh",
        [
            "Metode 1: Dropdown pada Baris Kode Verilog:",
            "  * Perhatikan kotak isian [ Blank_X v ] pada baris kode editor.",
            "  * Ketuk langsung kotak isian tersebut pada baris kode.",
            "  * Modal Bottom Sheet akan muncul dari bawah layar menampilkan pilihan opsi A, B, C, D.",
            "  * Ketuk opsi yang diinginkan -> Jawaban langsung terpasang pada kode.",
            "",
            "Fitur Cerdas Auto-Sizing (Teks Tampil Penuh):",
            "  * Kotak isian otomatis melebar menyesuaikan panjang kata yang dipilih.",
            "  * Kata pendek ('reg', 'b') maupun kata panjang ('assign', 'endmodule') selalu tampil 100% utuh tanpa huruf terpotong setengah.",
            "",
            "Metode 2: Kartu ChoiceChips di Bawah Kode:",
            "  * Mengelompokkan pilihan isian kode dalam kartu khusus.",
            "  * Sinkronisasi Dua Arah: Memilih chip seketika memperbarui kode di atasnya."
        ]
    )
    add_image_card(s4, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_code_blanks.png", "Pengisian Kode: Modal Bottom Sheet & Kotak Isian Auto-Sizing")

    # =========================================================================
    # SLIDE 5: Menjawab Pertanyaan Sub-Soal (Murni Pilihan Ganda Tanpa Text Box)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Menjawab Pertanyaan: Murni Pilihan Ganda (No Text Box)", "BAGIAN 4 DARI 5")
    add_card(s5, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Desain Murni Pilihan Ganda Responsif & Bebas Typo",
        [
            "Bebas Kotak Ketik Teks (No Text Box Clutter):",
            "  * Seluruh kotak input teks manual telah ditiadakan dari halaman soal.",
            "  * Layar tetap bersih dan bebas dari gangguan keyboard virtual smartphone yang menutupi diagram sirkuit.",
            "",
            "Menjawab dengan Satu Ketukan Sentuhan:",
            "  * Setiap pertanyaan logika/sinyal dilengkapi tombol pilihan (ChoiceChips).",
            "  * Cukup 1 sentuhan pada opsi jawaban yang diinginkan (misal: 'A. 0', 'B. 1', atau 'True/False').",
            "",
            "Tampilan Pilihan yang Dipilih (Selected Label):",
            "  * Di samping nomor Blank, label status otomatis menampilkan opsi yang dipilih:",
            "    'Selected: <Jawaban Anda>'.",
            "  * Teks jawaban yang dipilih terlihat penuh dan jelas, tidak terpotong.",
            "  * Bebas dari risiko salah ketik (typo) saat mengerjakan soal."
        ]
    )
    add_image_card(s5, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_questions.png", "Pertanyaan Sub-Soal: Murni Pilihan Ganda Tanpa Kotak Teks")

    # =========================================================================
    # SLIDE 6: Statistical Data: Evaluasi & Umpan Balik Real-Time
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Statistical Data: Evaluasi & Umpan Balik Real-Time", "BAGIAN 5 DARI 5")
    add_card(s6, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Ikhtisar Kartu Statistical Data pada Halaman Soal",
        [
            "Metrik Penilaian pada Statistical Data:",
            "  * Total number of questions : Jumlah seluruh titik isian (blanks) pada soal aktif.",
            "  * Number of correct answers : Jumlah isian yang dijawab dengan benar saat ini.",
            "  * Accuracy rate : Persentase ketepatan jawaban: (Correct / Total) x 100%.",
            "",
            "Indikator Warna Akurasi (Standar PLAS):",
            "  * Merah (< 60%) : Akurasi rendah, perlu kajian ulang materi.",
            "  * Netral (60% - 99%) : Sebagian besar benar, belum tuntas.",
            "  * Hijau Toska (100%) : Tuntas sempurna (Completed).",
            "",
            "Tombol Interaksi & Navigasi:",
            "  * 'Answer' : Menjalankan validasi SHA-256 instan dan memperbarui statistik.",
            "  * 'Previous' & 'Next' : Berpindah antar nomor latihan secara runtut.",
            "",
            "Kesimpulan:",
            "  * Statistical Data memberikan feedback terukur dan objektif secara mandiri!"
        ]
    )
    add_image_card(s6, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_statistical_data.png", "Tampilan Kartu Statistical Data & Kontrol Penilaian Real-Time")

    out_path = r"d:\APK_AI\VerilogPLAS_2\Tutorial_Mobile_VPLAS_Updated.pptx"
    prs.save(out_path)
    print(f"Updated 6-Slide Tutorial Presentation saved successfully to: {out_path}")

if __name__ == "__main__":
    build_presentation()
