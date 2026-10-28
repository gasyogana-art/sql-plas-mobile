import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    NAVY = RGBColor(0x2C, 0x3E, 0x50)       # #2C3E50
    DARK_BG = RGBColor(0x1A, 0x25, 0x2F)    # #1A252F
    TEAL = RGBColor(0x1B, 0xBC, 0x9B)       # #1BBC9B
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    LIGHT_GRAY = RGBColor(0xF8, 0xF9, 0xFA)
    BORDER_GRAY = RGBColor(0xDC, 0xE4, 0xEC)
    TEXT_DARK = RGBColor(0x2D, 0x37, 0x48)
    TEXT_MUTED = RGBColor(0x71, 0x80, 0x96)
    
    blank_layout = prs.slide_layouts[6]
    
    def add_header(slide, title_text, step_badge="TUTORIAL PENGGUNAAN"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.5), Inches(0.3))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = step_badge.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL
        
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.5), Inches(0.65))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY
        
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_GRAY
        line.line.color.rgb = BORDER_GRAY

    def add_card(slide, left, top, width, height, title, lines, title_color=NAVY, bg_color=LIGHT_GRAY, border_color=BORDER_GRAY):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)
        
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = title_color
        
        for line in lines:
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK
            p.space_before = Pt(4)
        return card

    def add_image_card(slide, left, top, width, height, image_path, caption=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_GRAY
        card.line.width = Pt(1)
        
        img_top = top + Inches(0.15)
        img_height = height - (Inches(0.6) if caption else Inches(0.3))
        img_left = left + Inches(0.15)
        img_width = width - Inches(0.3)
        
        if os.path.exists(image_path):
            slide.shapes.add_picture(image_path, img_left, img_top, img_width, img_height)
            
        if caption:
            cap_box = slide.shapes.add_textbox(left + Inches(0.2), top + height - Inches(0.45), width - Inches(0.4), Inches(0.35))
            tf_cap = cap_box.text_frame
            p_cap = tf_cap.paragraphs[0]
            p_cap.text = caption
            p_cap.font.size = Pt(10)
            p_cap.font.italic = True
            p_cap.font.color.rgb = TEXT_MUTED
            p_cap.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 1: Title Slide (TANPA GAMBAR / NO IMAGE - Sesuai Permintaan User)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BG
    bg1.line.fill.background()
    
    # Top Tag Badge
    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.9), Inches(3.6), Inches(0.36))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = TEAL
    badge1.line.fill.background()
    p_b = badge1.text_frame.paragraphs[0]
    p_b.text = "PANDUAN PRAKTIKUM MANDIRI"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = WHITE
    p_b.alignment = PP_ALIGN.CENTER
    
    # Big Title
    t_box1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(1.6))
    tf1 = t_box1.text_frame
    tf1.word_wrap = True
    p1_1 = tf1.paragraphs[0]
    p1_1.text = "Tutorial Penggunaan Mobile VPLAS Verilog"
    p1_1.font.size = Pt(36)
    p1_1.font.bold = True
    p1_1.font.color.rgb = WHITE
    
    p1_2 = tf1.add_paragraph()
    p1_2.text = "Panduan Praktis Pengerjaan Soal, Pengisian Kode, dan Evaluasi Mandiri"
    p1_2.font.size = Pt(18)
    p1_2.font.color.rgb = TEAL
    p1_2.space_before = Pt(4)
    
    # 2 Cards for 8 Tutorial Steps Overview (Slide 1 without image)
    # Card Left: Langkah 1 - 4
    card_l1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.2), Inches(5.6), Inches(3.7))
    card_l1.fill.solid()
    card_l1.fill.fore_color.rgb = NAVY
    card_l1.line.color.rgb = TEAL
    tf_cl1 = card_l1.text_frame
    tf_cl1.word_wrap = True
    p_cl1 = tf_cl1.paragraphs[0]
    p_cl1.text = "Alur Pengerjaan Soal (Bagian 1):"
    p_cl1.font.size = Pt(13)
    p_cl1.font.bold = True
    p_cl1.font.color.rgb = TEAL
    steps_left = [
        ("Langkah 1", "Memilih Soal pada Problems List & Status Badge"),
        ("Langkah 2", "Membaca Soal, Skema Sirkuit & Knowledge Point"),
        ("Langkah 3", "Mengisi Kode via Dropdown & Fitur Auto-Sizing"),
        ("Langkah 4", "Mengisi Kode via Kartu Pilihan Ganda (ChoiceChips)"),
    ]
    for stp, desc in steps_left:
        p = tf_cl1.add_paragraph()
        p.text = f"• [{stp}]  {desc}"
        p.font.size = Pt(12)
        p.font.color.rgb = WHITE
        p.space_before = Pt(10)

    # Card Right: Langkah 5 - 8
    card_r1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(3.2), Inches(5.6), Inches(3.7))
    card_r1.fill.solid()
    card_r1.fill.fore_color.rgb = NAVY
    card_r1.line.color.rgb = TEAL
    tf_cr1 = card_r1.text_frame
    tf_cr1.word_wrap = True
    p_cr1 = tf_cr1.paragraphs[0]
    p_cr1.text = "Alur Evaluasi & Ketuntasan (Bagian 2):"
    p_cr1.font.size = Pt(13)
    p_cr1.font.bold = True
    p_cr1.font.color.rgb = TEAL
    steps_right = [
        ("Langkah 5", "Menjawab Sub-Pertanyaan (Murni ChoiceChips)"),
        ("Langkah 6", "Evaluasi Jawaban Instan (Tombol Answer & Clear)"),
        ("Langkah 7", "Membaca Statistik Akurasi & Berpindah Soal"),
        ("Langkah 8", "Meninjau Ketuntasan Belajar pada Menu End"),
    ]
    for stp, desc in steps_right:
        p = tf_cr1.add_paragraph()
        p.text = f"• [{stp}]  {desc}"
        p.font.size = Pt(12)
        p.font.color.rgb = WHITE
        p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 2: Langkah 1 - Memilih Soal di Problems List
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Langkah 1: Memilih Soal pada Daftar Soal (Problems List)", "LANGKAH 1 DARI 8")
    add_card(s2, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Navigasi 22 Soal Bab 1 & Status Pengerjaan",
        [
            "Daftar 22 Soal Bab 1 Terstruktur:",
            "  * Soal 1-5  : Port wire, continuous assign, modul NOT gate.",
            "  * Soal 6-11 : Gerbang logika AND, OR, XOR, NAND, NOR.",
            "  * Soal 12-16: Multiplexer 2-to-1, Decoder, Adder, Subtractor.",
            "  * Soal 17-22: Rangkaian sekuensial Flip-Flop, Counter, & FSM.",
            "",
            "Mengenal Badge Status di Sisi Kanan Baris:",
            "  * Not Attempted (Tanpa Lencana): Soal belum pernah dicoba.",
            "  * Tried (Kuning): Pernah dikerjakan tetapi akurasi belum 100%.",
            "  * Completed (Hijau): Sempurna (seluruh isian benar 100%).",
            "",
            "Fitur Navigasi Tambahan:",
            "  * Kartu Progres Sesi: Menampilkan kemajuan pengerjaan X/22 secara langsung.",
            "  * Tombol 'TOP': Membantu menggulir cepat kembali ke posisi teratas.",
            "  * Ketuk baris soal untuk membuka layar pengerjaan soal."
        ]
    )
    add_image_card(s2, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_problems_list.png", "Tampilan Halaman Utama: Daftar Soal & Progres Pengerjaan")

    # =========================================================================
    # SLIDE 3: Langkah 2 - Membaca Soal, Diagram Sirkuit & Teori
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Langkah 2: Membaca Soal, Skema Sirkuit & Rangkuman Teori", "LANGKAH 2 DARI 8")
    add_card(s3, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Anatomi Layar Pengerjaan Soal",
        [
            "1. Circuit Diagrams & Waveforms (Posisi Paling Atas):",
            "   * Gambar skema sirkuit dan gelombang timing selalu muncul di posisi teratas layar.",
            "   * Fitur Zoom Interaktif: Gunakan gestur cubit (Pinch to Zoom) atau double-tap untuk memperbesar detail sinyal.",
            "",
            "2. Kartu Knowledge Point:",
            "   * Berisi teori ringkas dan kata kunci Verilog penting yang sedang diuji pada soal tersebut.",
            "   * Membantu mahasiswa memahami dasar teori sebelum mengisi kode.",
            "",
            "3. Editor Kode Verilog HDL:",
            "   * Menampilkan blok modul Verilog dengan nomor baris dan pewarnaan sintaks gelap yang nyaman di mata.",
            "   * Bagian rumpang ditandai dengan kontainer isian selektor."
        ]
    )
    add_image_card(s3, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "assets/images/Verilog_CH1_question2_img1.png", "Diagram Skema Rangkaian di Posisi Teratas Layar Soal")

    # =========================================================================
    # SLIDE 4: Langkah 3 - Mengisi Kode via Dropdown (Modal Pop-Up)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Langkah 3: Mengisi Kode via Dropdown & Fitur Auto-Sizing", "LANGKAH 3 DARI 8")
    add_card(s4, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Pengisian Cepat Langsung pada Baris Kode",
        [
            "Cara Membuka Pilihan Opsi:",
            "  1. Perhatikan kotak isian pada baris kode: [ Blank_1 v ].",
            "  2. Ketuk langsung kotak isian tersebut.",
            "  3. Jendela Modal Bottom Sheet akan muncul dari bawah layar menampilkan pilihan opsi A, B, C, D.",
            "  4. Ketuk opsi yang diinginkan -> Jendela menutup dan jawaban langsung terpasang.",
            "",
            "Fitur Cerdas Auto-Sizing (Teks Dijamin Tidak Terpotong):",
            "  * Kotak isian otomatis melebar menyesuaikan panjang kata kunci yang dipilih.",
            "  * Kata pendek ('reg', 'b') maupun kata panjang ('assign', 'endmodule') selalu tampil utuh dan jelas.",
            "  * Tidak ada huruf atau teks yang terpotong seperti pada versi web lama."
        ]
    )
    add_image_card(s4, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_code_blanks.png", "Demonstrasi Auto-Sizing & Modal Bottom Sheet Pilihan Ganda")

    # =========================================================================
    # SLIDE 5: Langkah 4 - Mengisi Kode via Kartu Pilihan Ganda
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Langkah 4: Mengisi Kode via Kartu Pilihan Ganda (ChoiceChips)", "LANGKAH 4 DARI 8")
    add_card(s5, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Alternatif Pengisian Efisien Tanpa Pop-Up",
        [
            "Kartu 'Multiple Choice: Code Blanks':",
            "  * Terletak tepat di bawah blok kode Verilog.",
            "  * Mengelompokkan pilihan ganda untuk setiap Blank [1], Blank [2], dst. dalam satu tampilan kartu.",
            "",
            "Cara Menjawab Cukup 1 Sentuhan:",
            "  * Cukup ketuk tombol chip pilihan (A, B, C, atau D).",
            "  * Chip yang dipilih langsung berubah warna menjadi Hijau Teal.",
            "  * Label status di atas chip menampilkan: 'Selected: <kata_kunci>'.",
            "",
            "Sinkronisasi Otomatis Dua Arah (2-Way Sync):",
            "  * Memilih di kartu ChoiceChips seketika memperbarui tampilan baris kode di atasnya.",
            "  * Begitu pula jika memilih lewat modal dropdown, kartu ini ikut tersinkronisasi."
        ]
    )
    add_image_card(s5, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_code_chips.png", "Kartu Pilihan Ganda Code Blanks dengan Sinkronisasi Dua Arah")

    # =========================================================================
    # SLIDE 6: Langkah 5 - Menjawab Sub-Pertanyaan (Murni Pilihan Ganda)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Langkah 5: Menjawab Pertanyaan Sub-Soal (Tanpa Text Box)", "LANGKAH 5 DARI 8")
    add_card(s6, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Desain Murni Pilihan Ganda Responsif & Bebas Typo",
        [
            "Bebas Kotak Ketik Teks (No Text Box Clutter):",
            "  * Seluruh kotak input teks manual telah dihilangkan.",
            "  * Mahasiswa tidak perlu lagi mengetik di keyboard HP yang rentan typo (salah ketik) atau menutupi diagram sirkuit.",
            "",
            "Menjawab dengan Satu Ketukan:",
            "  * Setiap pertanyaan teori/sinyal dilengkapi tombol pilihan (ChoiceChips).",
            "  * Ketuk opsi yang diinginkan (misal: 'A. 0', 'B. 1', atau 'True/False').",
            "  * Opsi terpilih langsung ditandai dengan warna teal solid.",
            "",
            "Kemudahan Mengubah Jawaban:",
            "  * Cukup ketuk opsi lain jika ingin mengganti pilihan sebelum dievaluasi.",
            "  * Sangat cepat, ergonomis, dan nyaman untuk penggunaan satu tangan."
        ]
    )
    add_image_card(s6, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_questions.png", "Antarmuka Pertanyaan Murni Pilihan Ganda Bebas Kotak Teks")

    # =========================================================================
    # SLIDE 7: Langkah 6 - Melakukan Evaluasi (Tombol Answer & Clear)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Langkah 6: Evaluasi Jawaban & Indikator Hasil Koreksi", "LANGKAH 6 DARI 8")
    add_card(s7, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Validasi Instan Rumus Hash SHA-256 Asli PLAS",
        [
            "Menjalankan Evaluasi dengan Tombol 'Answer':",
            "  * Setelah semua isian kode dan pertanyaan dipilih, ketuk tombol 'Answer' di bawah.",
            "  * Sistem melakukan verifikasi enkripsi instan:",
            "    Hash = SHA256( jawaban + statementId + index ).",
            "",
            "Umpan Balik Warna Langsung di Layar:",
            "  * Jawaban Benar : Berubah HIJAU TEAL dengan ikon centang [✓].",
            "  * Jawaban Salah : Berubah MERAH dengan ikon silang [✗].",
            "",
            "Tombol 'Clear Answers' (Reset Jawaban):",
            "  * Ikon refresh di AppBar atas untuk mengosongkan seluruh pilihan pada soal aktif.",
            "  * Gunakan fitur ini jika ingin mencoba kembali soal dari awal."
        ]
    )
    add_image_card(s7, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_evaluation.png", "Tampilan Evaluasi Jawaban Tombol Answer (Centang Hijau & Silang Merah)")

    # =========================================================================
    # SLIDE 8: Langkah 7 - Membaca Skor & Standar Akurasi
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Langkah 7: Membaca Statistik Akurasi & Berpindah Soal", "LANGKAH 7 DARI 8")
    add_card(s8, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Kartu Statistical Data & Standar Kelulusan",
        [
            "Metrik pada Kartu 'Statistical Data':",
            "  * Questions : Total jumlah isian kode dan pertanyaan pada soal.",
            "  * Correct   : Jumlah isian yang berhasil dijawab benar.",
            "  * Accuracy  : Nilai persentase akurasi pengerjaan saat ini.",
            "",
            "Standar Pewarnaan Akurasi Resmi PLAS:",
            "  * < 60.00%        : Teks MERAH (Belum memenuhi syarat kelulusan).",
            "  * 60.00% - 99.99% : Teks GELAP (Lulus sebagian, perlu perbaikan).",
            "  * 100.00%         : Teks HIJAU TEAL (Tuntas sempurna / All Correct).",
            "",
            "Navigasi Berpindah Soal:",
            "  * Tombol '< Previous Problem' : Kembali ke nomor soal sebelumnya.",
            "  * Tombol 'Next Problem >'     : Melompat ke nomor soal berikutnya."
        ]
    )
    add_image_card(s8, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_stats.png", "Kartu Statistical Data & Standar Nilai Akurasi")

    # =========================================================================
    # SLIDE 9: Langkah 8 - Meninjau Ketuntasan di Tab End
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Langkah 8: Meninjau Ketuntasan Belajar pada Menu End", "LANGKAH 8 DARI 8")
    add_card(s9, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.3),
        "Ikhtisar Akhir Sesi & Tombol Kontrol Praktikum",
        [
            "Ringkasan Sesi (Session Overview):",
            "  * Completed (Hijau) : Jumlah soal tuntas 100% (target: 22 soal).",
            "  * Tried (Kuning)    : Soal yang masih perlu diperbaiki.",
            "  * Pending (Abu-abu) : Soal yang belum dikerjakan.",
            "",
            "Bilah Progres Total Bab 1:",
            "  * Menampilkan persentase ketuntasan keseluruhan secara grafis visual (misal: 100% Selesai).",
            "",
            "Aksi Kontrol Praktikum:",
            "  * 'Continue Solving Problems': Melompat kembali ke Problems List.",
            "  * 'Reset All Local Records': Mengosongkan database lokal dengan konfirmasi pengaman jika ingin berlatih ulang.",
            "  * 'About PLAS': Informasi versi sistem dan lisensi riset.",
            "",
            "Kesimpulan:",
            "  * Mahasiswa dapat belajar dan berlatih mandiri secara praktis, cepat, dan terukur!"
        ]
    )
    add_image_card(s9, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.3), "scratch/slides_media/ui_end_screen.png", "Menu End: Ringkasan Sesi Ketuntasan Belajar & Kontrol Data")

    out_path = r"d:\APK_AI\VerilogPLAS_2\Tutorial_Mobile_VPLAS.pptx"
    prs.save(out_path)
    print(f"Focused Tutorial Presentation saved successfully to: {out_path}")

if __name__ == "__main__":
    build_presentation()
